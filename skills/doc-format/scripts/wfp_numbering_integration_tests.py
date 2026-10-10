"""Platform/configuration integration tests, without launching Office."""
import contextlib
import io
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock, patch
from zipfile import BadZipFile

from docx import Document
from docx.oxml.ns import qn

from wfp_config import DEFAULT_CONFIG
from wfp_core import WordProcessor
from wfp_numbering_tests import add_list, element, numbered, paragraphs, read_parts, texts, NS


class NumberingIntegrationTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='wfp_numbering_')
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.source = self.root / 'input.docx'
        self.output = self.root / 'output.docx'
        self.logs = []
        self.processor = WordProcessor(DEFAULT_CONFIG.copy(), self.logs.append)
        self.addCleanup(self.processor._cleanup_temp_files)

    def fixture(self, unsupported=False):
        doc = Document()
        doc.add_paragraph('报告标题')
        add_list(doc,100,[{'fmt':'japaneseCounting','text':'%1、','suffix':'tab'}])
        add_list(doc,101,[{'text':'%1.','suffix':'tab'}])
        numbered(doc.add_paragraph('标题正文'),100,0)
        doc.add_paragraph('二、手写标题')
        numbered(doc.add_paragraph('测试'),101,0)
        numbered(doc.add_paragraph('好的'),101,0)
        numbered(doc.add_paragraph('  原文空格'),101,0)
        numbered(doc.add_table(rows=1,cols=1).cell(0,0).paragraphs[0],101,0).add_run('表格正文')
        if unsupported:
            add_list(doc,102,[{'fmt':'cardinalText','text':'%1.'}])
            numbered(doc.add_paragraph('不支持的编号'),102,0)
        doc.save(self.source)

    def test_non_windows_default_preprocessing_preserves_source_and_layout(self):
        self.fixture()
        original = read_parts(self.source)
        self.processor.config.pop('convert_auto_numbering')  # older config
        with patch('wfp_core.IS_WINDOWS',False), patch.object(self.processor,'_preprocess_com_tasks') as com:
            output = self.processor._preprocess_document(str(self.source))
        com.assert_not_called()
        self.assertEqual(texts(output)[1:7],['一、标题正文','二、手写标题','1.测试','2.好的','3.  原文空格','4.表格正文'])
        self.assertEqual(read_parts(self.source),original)
        for name,data in read_parts(output).items():
            if name != 'word/document.xml': self.assertEqual(data,original[name])
        ind = paragraphs(output)[3].find('w:pPr/w:ind',NS)
        self.assertEqual(ind.get(qn('w:left')),'720')
        self.assertEqual(ind.get(qn('w:hanging')),'360')
        self.processor._cleanup_temp_files()
        self.assertFalse(Path(output).exists())

    def test_disabled_keeps_automatic_numbering_and_does_not_invoke_converter(self):
        self.fixture()
        self.processor.config['convert_auto_numbering'] = False
        with patch('wfp_core.IS_WINDOWS',False), patch('wfp_numbering.convert_docx_numbering') as convert:
            self.processor.format_document(str(self.source),str(self.output))
        convert.assert_not_called()
        self.assertEqual(paragraphs(self.output)[1].find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'100')
        self.assertNotIn('一、标题正文',texts(self.output))

    def test_conversion_precedes_heading_detection_and_includes_unformatted_table(self):
        self.fixture()
        self.processor.config['enable_table_formatting'] = False
        with patch('wfp_core.IS_WINDOWS',False):
            self.processor.format_document(str(self.source),str(self.output))
        result = paragraphs(self.output)
        self.assertEqual(result[1].find('w:pPr/w:outlineLvl',NS).get(qn('w:val')),'0')
        self.assertIn('1.测试',texts(self.output))
        self.assertIn('2.好的',texts(self.output))
        self.assertIn('4.表格正文',texts(self.output))
        for i in (1,3,4,5,6):
            self.assertEqual(result[i].find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'0')
        paths = list(self.processor.temp_files)
        self.assertEqual(len(paths),2)
        self.processor._cleanup_temp_files()
        self.assertTrue(all(not Path(p).exists() for p in paths))

    def test_partial_conversion_keeps_unsupported_numbering_and_logs_warning(self):
        self.fixture(unsupported=True)
        with patch('wfp_core.IS_WINDOWS',False):
            self.processor.format_document(str(self.source),str(self.output))
        self.assertIn('1.测试',texts(self.output))
        self.assertEqual(paragraphs(self.output)[-1].find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'102')
        self.assertTrue(any('部分转换' in s for s in self.logs))
        self.assertTrue(any('cardinalText' in s for s in self.logs))
        Document(self.output)

    def test_numbering_revisions_are_retained_in_partial_mode(self):
        self.fixture()
        doc=Document(self.source)
        ppr=doc.paragraphs[1]._p.get_or_add_pPr()
        ppr.append(element('pPrChange',id=1))
        doc.save(self.source)
        with patch('wfp_core.IS_WINDOWS',False):
            output=self.processor._preprocess_document(str(self.source))
        p=paragraphs(output)[1]
        self.assertIsNotNone(p.find('w:pPr/w:pPrChange',NS))
        self.assertEqual(p.find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'100')
        self.assertIn('1.测试',texts(output))
        self.assertTrue(any('修订' in s for s in self.logs))

    def test_windows_keeps_com_first_and_falls_back_when_unavailable(self):
        self.fixture()
        for enabled in (False,True):
            for available in (False,True):
                with self.subTest(enabled=enabled,available=available):
                    self.processor.config['convert_auto_numbering']=enabled
                    app=MagicMock()
                    app.Documents.Open.return_value.Revisions.Count=0
                    with patch('wfp_core.IS_WINDOWS',True), \
                         patch.object(self.processor,'_com_available',return_value=available), \
                         patch.object(self.processor,'_get_wps_app',return_value=app), \
                         patch('wfp_numbering.convert_docx_numbering') as convert:
                        convert.return_value.converted=5
                        convert.return_value.warnings=[]
                        path=self.processor._preprocess_document(str(self.source))
                    if available:
                        self.assertEqual(path,str(self.source))
                        convert.assert_not_called()
                    else:
                        self.assertNotEqual(path,str(self.source))
                        convert.assert_called_once_with(str(self.source),path,strict=False,preserve_suffix=False)
                    self.assertEqual(app.Documents.Open.call_count,int(available))
                    self.assertEqual(app.Documents.Open.return_value.Content.ListFormat.ConvertNumbersToText.call_count,int(available))

    def test_windows_office_not_installed_uses_python_fallback_end_to_end(self):
        self.fixture()
        # pywin32 can be installed even when neither Office application is.
        with patch('wfp_core.IS_WINDOWS',True), \
             patch.object(self.processor,'_com_available',return_value=True), \
             patch.object(self.processor,'_get_wps_app',side_effect=RuntimeError('Office not installed')):
            self.processor.format_document(str(self.source),str(self.output))
        self.assertIn('1.测试',texts(self.output))
        self.assertIn('4.表格正文',texts(self.output))
        self.assertTrue(any('回退到 Python' in s for s in self.logs))

    def test_legacy_conversion_feeds_numbering_preprocessing(self):
        self.fixture()
        for ext in ('.doc','.wps'):
            legacy=self.root/('legacy'+ext)
            legacy.write_bytes(b'legacy fixture handled by mock')
            def converted(_source,destination):
                shutil.copy2(self.source,destination)
            with patch('wfp_core.IS_WINDOWS',False), \
                 patch.object(self.processor,'_convert_legacy_with_soffice',side_effect=converted) as soffice:
                self.processor.format_document(str(legacy),str(self.output))
            soffice.assert_called_once()
            self.assertIn('1.测试',texts(self.output))
            self.processor._cleanup_temp_files()

    def test_text_and_markdown_skip_numbering(self):
        for ext in ('.txt','.md'):
            source=self.root/('text'+ext)
            source.write_text('报告标题\n1. 手写正文',encoding='utf-8')
            with patch('wfp_core.IS_WINDOWS',False), patch('wfp_numbering.convert_docx_numbering') as convert:
                self.processor.format_document(str(source),str(self.output))
            convert.assert_not_called()
            self.assertTrue(self.output.exists())
            self.processor._cleanup_temp_files()

    def test_fatal_read_error_does_not_write_output_and_temp_files_are_cleaned(self):
        self.source.write_bytes(b'not a zip')
        with patch('wfp_core.IS_WINDOWS',False), self.assertRaises(BadZipFile):
            self.processor.format_document(str(self.source),str(self.output))
        paths=list(self.processor.temp_files)
        self.processor._cleanup_temp_files()
        self.assertTrue(all(not Path(p).exists() for p in paths))
        self.assertFalse(self.output.exists())
        self.assertEqual(self.source.read_bytes(),b'not a zip')

    def test_cli_config_default_disable_save_and_reload(self):
        import wfp_cli
        self.assertTrue(wfp_cli.normalize_config({})['convert_auto_numbering'])
        path=self.root/'config.json'
        args=wfp_cli.build_parser().parse_args(['save-config','--set','convert_auto_numbering=false','-o',str(path)])
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(args.func(args),0)
        config,_=wfp_cli.load_config(config_path=path)
        self.assertIs(config['convert_auto_numbering'],False)
        self.assertIs(json.loads(path.read_text(encoding='utf-8'))['convert_auto_numbering'],False)


if __name__=='__main__':
    unittest.main()
