"""Regression fixtures for the standalone converter; no formatter imports."""
import base64
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from wfp_numbering import W, NumberingConversionError, convert_docx_numbering, format_number

NS = {'w': W}
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
OFFICE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'


def element(local_name, **attrs):
    result = OxmlElement('w:' + local_name)
    for key, value in attrs.items():
        result.set(qn('w:' + key), str(value))
    return result


def numbered(paragraph, number, level=None):
    ppr = paragraph._p.get_or_add_pPr() if hasattr(paragraph, '_p') else paragraph.find(qn('w:pPr'))
    if ppr is None:
        ppr = element('pPr')
        paragraph.insert(0, ppr)
    old = ppr.find(qn('w:numPr'))
    if old is not None:
        ppr.remove(old)
    numpr = element('numPr')
    if level is not None:
        numpr.append(element('ilvl', val=level))
    numpr.append(element('numId', val=number))
    ppr.append(numpr)
    return paragraph


def add_list(document, number=100, levels=None, abstract_id=None, overrides=None):
    """Add an explicit, deterministic list unrelated to python-docx defaults."""
    abstract_id = number if abstract_id is None else abstract_id
    levels = levels or [{'fmt': 'decimal', 'text': '%1.'}]
    root = document.part.numbering_part.element
    abstract = element('abstractNum', abstractNumId=abstract_id)
    abstract.append(element('multiLevelType', val='multilevel'))
    for index, properties in enumerate(levels):
        lvl = element('lvl', ilvl=index)
        lvl.append(element('start', val=properties.get('start', 1)))
        lvl.append(element('numFmt', val=properties.get('fmt', 'decimal')))
        if 'restart' in properties:
            lvl.append(element('lvlRestart', val=properties['restart']))
        if 'style' in properties:
            lvl.append(element('pStyle', val=properties['style']))
        if properties.get('legal'):
            lvl.append(element('isLgl'))
        lvl.append(element('suff', val=properties.get('suffix', 'nothing')))
        lvl.append(element('lvlText', val=properties.get('text', f'%{index+1}.')))
        lvl.append(element('lvlJc', val='left'))
        ppr = element('pPr')
        tabs = element('tabs')
        tabs.append(element('tab', val='num', pos=720 * (index+1)))
        ppr.append(tabs)
        ppr.append(element('ind', left=720 * (index+1), hanging=360))
        lvl.append(ppr)
        rpr = element('rPr')
        rpr.append(element('rFonts', ascii='Arial', eastAsia='宋体'))
        lvl.append(rpr)
        abstract.append(lvl)
    # abstractNum elements precede num instances in the schema.
    root.insert(0, abstract)
    instance = element('num', numId=number)
    instance.append(element('abstractNumId', val=abstract_id))
    for level, start in (overrides or {}).items():
        override = element('lvlOverride', ilvl=level)
        override.append(element('startOverride', val=start))
        instance.append(override)
    root.append(instance)
    return abstract, instance


def read_parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rewrite_package(path, updates):
    parts = read_parts(path)
    parts.update(updates)
    with ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def paragraphs(path, part='word/document.xml'):
    return list(etree.fromstring(read_parts(path)[part]).iter(qn('w:p')))


def text(paragraph):
    # Include tabs, but only this paragraph's runs (not nested textboxes).
    result = []
    for item in paragraph.iter():
        owner = next(item.iterancestors(qn('w:p')), None)
        if owner is not paragraph:
            continue
        if any(a.tag == qn('w:pPr') for a in item.iterancestors()):
            continue
        if item.tag == qn('w:t'):
            result.append(item.text or '')
        elif item.tag == qn('w:tab'):
            result.append('\t')
    return ''.join(result)


def texts(path, part='word/document.xml'):
    return [text(p) for p in paragraphs(path, part)]


def mixed_fixture(path):
    document = Document()
    add_list(document, 100, [{'fmt': 'chineseCounting', 'text': '%1、'}])
    add_list(document, 101, [{'fmt': 'chineseCounting', 'text': '（%1）', 'start': 2}])
    add_list(document, 102)
    numbered(document.add_paragraph('自动一级'), 100, 0)
    document.add_paragraph('二、手写一级')
    document.add_paragraph('（一）手写二级')
    numbered(document.add_paragraph('自动二级'), 101, 0)
    numbered(document.add_paragraph('自动三级'), 102, 0)
    document.add_paragraph('2.手写三级')
    table = document.add_table(rows=1, cols=2)
    numbered(table.cell(0, 0).paragraphs[0], 102, 0).add_run('表格内自动编号')
    table.cell(0, 1).text = '3.表格内手写编号'
    numbered(document.add_paragraph('表格后自动编号'), 102, 0)
    document.save(path)


def oracle_fixture(path):
    """Extra cases verified with WPS, not required for running unit tests."""
    document = Document()
    formats = ['upperLetter', 'lowerLetter', 'upperRoman', 'decimalZero',
               'chineseCounting', 'chineseCountingThousand', 'chineseLegalSimplified',
               'ideographTraditional', 'ideographZodiac', 'decimalEnclosedCircle',
               'decimalEnclosedParen', 'decimalEnclosedFullstop']
    number = 100
    for fmt in formats:
        for value in (0, 1, 10, 11, 20, 26, 27, 28, 52, 53, 101, 1001, 10001):
            if fmt in ('upperLetter','lowerLetter') and value == 0:
                continue
            if fmt == 'upperRoman' and not 1 <= value <= 999:
                continue
            add_list(document, number, [{'fmt': fmt, 'text': '%1|', 'start': value}])
            numbered(document.add_paragraph(f'{fmt}:{value}'), number, 0)
            number += 1
    for restart in (None, 0, 1, 2, 3):
        levels = [{'text': '%1.'}, {'text': '%1.%2.'}, {'text': '%1.%2.%3.'}]
        if restart is not None:
            levels[2]['restart'] = restart
        add_list(document, number, levels)
        for index, level in enumerate((0, 1, 2, 2, 0, 2, 1, 2)):
            numbered(document.add_paragraph(f'restart={restart};step={index};level={level}'), number, level)
        number += 1
    document.save(path)


def stories_fixture(path):
    """Valid package containing notes referenced by the main story."""
    document = Document()
    add_list(document)
    numbered(document.add_paragraph('正文A'),100,0)
    numbered(document.sections[0].header.paragraphs[0],100,0).add_run('页眉')
    numbered(document.sections[0].footer.paragraphs[0],100,0).add_run('页脚')
    for singular in ('footnote','endnote'):
        for identity in (1,2):
            p = document.add_paragraph(singular+str(identity))
            p.add_run()._r.append(element(singular+'Reference',id=identity))
    numbered(document.add_paragraph('正文B'),100,0)
    document.save(path)
    parts = read_parts(path)
    types = etree.fromstring(parts['[Content_Types].xml'])
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    updates = {}
    for name,singular in [('footnotes','footnote'),('endnotes','endnote')]:
        root = element(name)
        for identity in (1,2):
            note = element(singular,id=identity)
            for suffix in ('A','B'):
                p = element('p'); numbered(p,100,0)
                r = element('r'); t = element('t'); t.text = singular+str(identity)+suffix
                r.append(t); p.append(r); note.append(p)
            root.append(note)
        updates['word/'+name+'.xml'] = etree.tostring(root)
        etree.SubElement(types,'{'+CT+'}Override',PartName='/word/'+name+'.xml',
            ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.'+name+'+xml')
        etree.SubElement(rels,'{'+REL+'}Relationship',Id='rId'+name,Type=OFFICE_REL+name,Target=name+'.xml')
    updates['[Content_Types].xml'] = etree.tostring(types)
    updates['word/_rels/document.xml.rels'] = etree.tostring(rels)
    rewrite_package(path,updates)


def formatting_fixture(path):
    document = Document()
    add_list(document,100,[{'fmt':'chineseCounting','text':'%1、','suffix':'tab'}])
    base = document.styles.add_style('IndentedList',WD_STYLE_TYPE.PARAGRAPH)
    ppr = element('pPr')
    ppr.append(element('ind',left=1800,hanging=600))
    base.element.append(ppr)
    numbered(base.element,100)
    for suffix in ('A','B','C'):
        p = document.add_paragraph(suffix,style=base)
        if suffix == 'B':
            p._p.get_or_add_pPr().append(element('ind',left=1000))
        if suffix == 'C':
            p._p.get_or_add_pPr().append(element('ind',hanging=100))
    # Bound style inheritance with a direct level override and numId=0.
    child = document.styles.add_style('BoundChild',WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    child.element.append(element('pPr'))
    numbered(child.element,100,1)
    child.element.find(qn('w:pPr')).append(element('ind',left=1800,hanging=600))
    abstract = document.part.numbering_part.element.find("w:abstractNum[@w:abstractNumId='100']",NS)
    lvl = deepcopy(abstract.find(qn('w:lvl')))
    lvl.set(qn('w:ilvl'),'1')
    lvl.find(qn('w:lvlText')).set(qn('w:val'),'（%2）')
    lvl.insert(2,element('pStyle',val='BoundChild'))
    abstract.append(lvl)
    document.add_paragraph('绑定子样式',style=child)
    numbered(document.add_paragraph('手写',style=child),0)
    p = document.add_paragraph('直接指定第一级',style=child)
    np = element('numPr'); np.append(element('ilvl',val=0)); p._p.get_or_add_pPr().append(np)
    document.save(path)


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / 'input.docx'
        self.output = Path(self.directory.name) / 'output.docx'

    def convert(self, document, **options):
        document.save(self.source)
        return convert_docx_numbering(self.source, self.output, **options)

    def boundary_document(self, *, suffix='tab', hanging=420, custom=None):
        document = Document()
        abstract,_ = add_list(document,levels=[{'fmt':'japaneseCounting','text':'%1、','suffix':suffix}])
        lvl = abstract.find('w:lvl',NS)
        ppr = lvl.find('w:pPr',NS)
        tabs = ppr.find('w:tabs',NS)
        if custom is None:
            ppr.remove(tabs)
        else:
            tabs[0].set(qn('w:pos'),str(custom))
        ind = ppr.find('w:ind',NS)
        ind.set(qn('w:left'),'840'); ind.set(qn('w:hanging'),str(hanging))
        lvl.find('w:rPr',NS).append(element('sz',val=21))
        return document,lvl

    def test_cjk_boundary_tab_in_body_table_and_manual_spaces(self):
        document,lvl = self.boundary_document()
        # Word often omits suff: its defined default is still a tab.
        lvl.remove(lvl.find('w:suff',NS))
        numbered(document.add_paragraph('  原文空格保留'),100,0)
        numbered(document.add_table(rows=1,cols=1).cell(0,0).paragraphs[0],100,0).add_run('表格正文')
        document.add_paragraph('三、  手写编号')
        self.convert(document,preserve_suffix=True)
        self.assertEqual(texts(self.output),['一、  原文空格保留','二、表格正文','三、  手写编号'])
        for p in paragraphs(self.output)[:2]:
            self.assertEqual(p.find('w:pPr/w:ind',NS).get(qn('w:hanging')),'420')

    def test_cjk_number_growth_and_required_tabs(self):
        document,_ = self.boundary_document()
        for _ in range(11):
            numbered(document.add_paragraph('正文'),100,0)
        self.convert(document,preserve_suffix=True)
        self.assertEqual(texts(self.output)[0],'一、正文')
        self.assertEqual(texts(self.output)[9],'十、正文')
        self.assertEqual(texts(self.output)[10],'十一、\t正文')

    def test_cjk_separator_guards(self):
        for variant in ('short','long','space','nothing','custom_after','font_spacing',
                        'font_scale','superscript','right_aligned','character_indent',
                        'compat','character_grid','unknown_size','earlier_tab'):
            with self.subTest(variant=variant):
                document,lvl = self.boundary_document(
                    suffix=variant if variant in ('space','nothing') else 'tab',
                    hanging={'short':520,'long':320}.get(variant,420),
                    custom={'custom_after':1200,'earlier_tab':600}.get(variant))
                rpr=lvl.find('w:rPr',NS)
                if variant=='font_spacing': rpr.append(element('spacing',val=1))
                if variant=='font_scale': rpr.append(element('w',val=90))
                if variant=='superscript': rpr.append(element('vertAlign',val='superscript'))
                if variant=='unknown_size':
                    rpr.remove(rpr.find('w:sz',NS))
                    defaults=document.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr',NS)
                    for node in defaults.findall('w:sz',NS): defaults.remove(node)
                if variant=='right_aligned': lvl.find('w:lvlJc',NS).set(qn('w:val'),'right')
                if variant=='character_indent': lvl.find('w:pPr/w:ind',NS).set(qn('w:hangingChars'),'200')
                if variant=='compat':
                    document.settings.element.find('w:compat',NS).append(element('doNotUseIndentAsNumberingTabStop'))
                if variant=='character_grid':
                    document.sections[0]._sectPr.append(element('docGrid',type='linesAndChars',charSpace=100))
                numbered(document.add_paragraph('正文'),100,0)
                document.save(self.source)
                convert_docx_numbering(self.source,self.output,overwrite=True,preserve_suffix=True)
                separator = ' ' if variant=='space' else '' if variant in ('nothing','earlier_tab') else '\t'
                self.assertEqual(texts(self.output),['一、'+separator+'正文'])


    def test_default_omits_only_generated_suffix_in_body_and_table(self):
        document=Document()
        originals=[]
        for i,suffix in enumerate(('tab','space','nothing')):
            add_list(document,100+i,[{'text':'%1.','suffix':suffix}])
            for body in ('正文','  原文空格','\t原文制表符'):
                numbered(document.add_paragraph(body),100+i,0)
                originals.append(body)
            p=document.add_table(rows=1,cols=1).cell(0,0).paragraphs[0]
            numbered(p,100+i,0).add_run('表格正文')
            originals.append('表格正文')
        document.add_paragraph('5.  手写编号')
        originals.append('5.  手写编号')
        self.convert(document)
        expected=[f'{i%4+1}.'+body for i,body in enumerate(originals[:-1])]+[originals[-1]]
        self.assertEqual(texts(self.output),expected)
        # Adding no suffix must not clear the original hanging/left indent.
        for p in paragraphs(self.output)[:-1]:
            ind=p.find('w:pPr/w:ind',NS)
            self.assertEqual(ind.get(qn('w:left')),'720')
            self.assertEqual(ind.get(qn('w:hanging')),'360')
        again=self.output.with_name('again.docx')
        self.assertEqual(convert_docx_numbering(self.output,again).converted,0)
        self.assertEqual(read_parts(again),read_parts(self.output))

    def test_preserve_suffix_is_explicit_opt_in(self):
        document=Document()
        add_list(document,100,[{'text':'%1.','suffix':'tab'}])
        add_list(document,101,[{'text':'%1.','suffix':'space'}])
        numbered(document.add_paragraph('  正文'),100,0)
        numbered(document.add_paragraph('\t正文'),101,0)
        self.convert(document,preserve_suffix=True)
        self.assertEqual(texts(self.output),['1.\t  正文','1. \t正文'])

    def test_mixed_manual_and_automatic_including_table(self):
        mixed_fixture(self.source)
        report = convert_docx_numbering(self.source, self.output)
        self.assertEqual(report.converted, 5)
        self.assertEqual(texts(self.output), ['一、自动一级', '二、手写一级', '（一）手写二级',
            '（二）自动二级', '1.自动三级', '2.手写三级', '2.表格内自动编号',
            '3.表格内手写编号', '3.表格后自动编号'])

    def test_multilevel_restart_placeholders_start_override_and_legal(self):
        document = Document()
        add_list(document, levels=[{'fmt':'upperRoman','text':'%1.'},
            {'fmt':'lowerLetter','text':'%1.%2.'}, {'text':'%1.%2.%3.','legal':True}])
        for level in (0, 1, 2, 2, 1, 2, 0, 1, 2):
            numbered(document.add_paragraph('X'), 100, level)
        _, instance = add_list(document, 101, overrides={0:5})
        instance.find(qn('w:abstractNumId')).set(qn('w:val'), '100')
        for _ in range(2):
            numbered(document.add_paragraph('Y'), 101, 0)
        self.convert(document)
        self.assertEqual(texts(self.output), ['I.X','I.a.X','1.1.1.X','1.1.2.X',
            'I.b.X','1.2.1.X','II.X','II.a.X','2.1.1.X','V.Y','VI.Y'])

    def test_never_restart_and_explicit_restart(self):
        document = Document()
        for number, restart in ((100, 0), (101, 1)):
            add_list(document, number, [{'text':'%1.'}, {'text':'%2.'},
                {'text':'%3.', 'restart':restart}])
            for level in (0, 1, 2, 2, 1, 2, 0, 1, 2):
                numbered(document.add_paragraph('X'), number, level)
        self.convert(document)
        self.assertEqual(texts(self.output), ['1.X','1.X','1.X','2.X','2.X','3.X',
            '2.X','1.X','4.X', '1.X','1.X','1.X','2.X','2.X','3.X','2.X','1.X','1.X'])

    def test_style_inheritance_binding_partial_numpr_and_suppression(self):
        document = Document()
        base = document.styles.add_style('NumberBase', WD_STYLE_TYPE.PARAGRAPH)
        child = document.styles.add_style('NumberChild', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        child.element.append(element('pPr'))
        numbered(child.element,100,1)
        base.element.append(element('pPr'))
        numbered(base.element, 100)  # accepts any XML node with pPr
        add_list(document, levels=[{'text':'%1.'}, {'text':'（%2）', 'style':'NumberChild'}])
        document.add_paragraph('A', style=child)
        suppressed = document.add_paragraph('手写', style=child)
        numbered(suppressed, 0)
        partial = document.add_paragraph('B', style=child)
        np = element('numPr')
        np.append(element('ilvl', val=0))
        partial._p.get_or_add_pPr().append(np)
        document.add_paragraph('C', style=child)
        self.convert(document)
        self.assertEqual(texts(self.output), ['（1）A','手写','2.B','（1）C'])

    def test_linked_numbering_style_and_complete_level_override(self):
        document = Document()
        add_list(document, 100)
        style = document.styles.add_style('NumberLink', WD_STYLE_TYPE.LIST)
        style.element.append(element('pPr'))
        numbered(style.element, 100)
        abstract, instance = add_list(document, 101)
        abstract[:] = [element('numStyleLink', val='NumberLink')]
        override = element('lvlOverride', ilvl=0)
        lvl = element('lvl', ilvl=0)
        for name, value in [('start',3),('numFmt','chineseCounting'),('suff','space'),('lvlText','第%1项')]:
            lvl.append(element(name,val=value))
        override.append(lvl)
        instance.append(override)
        for word in ('A', 'B'):
            numbered(document.add_paragraph(word), 101, 0)
        self.convert(document)
        self.assertEqual(texts(self.output), ['第三项A','第四项B'])

    def test_style_level_without_pstyle_binding(self):
        document = Document()
        add_list(document,levels=[{'text':'%1.'},{'text':'（%2）','fmt':'chineseCounting'}])
        style = document.styles.add_style('LevelOnly',WD_STYLE_TYPE.PARAGRAPH)
        style.element.append(element('pPr'))
        numbered(style.element,100,1)
        document.add_paragraph('A',style=style)
        document.add_paragraph('B',style=style)
        self.convert(document)
        self.assertEqual(texts(self.output),['（一）A','（二）B'])

    def test_images_tables_bookmarks_hyperlinks_and_package_contents_preserved(self):
        document = Document()
        add_list(document, levels=[{'text':'%1.', 'suffix':'tab'}])
        p = numbered(document.add_paragraph('保留粗体'), 100, 0)
        p.runs[0].bold = True
        p._p.insert(1, element('bookmarkStart', id=12, name='keep'))
        p._p.append(element('bookmarkEnd', id=12))
        hyperlink = element('hyperlink', anchor='keep')
        run = element('r')
        t = element('t')
        t.text = '原有链接'
        run.append(t)
        hyperlink.append(run)
        p._p.append(hyperlink)
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII=')
        p.add_run().add_picture(BytesIO(png))
        table = document.add_table(rows=2, cols=2)
        table.cell(0,0).merge(table.cell(0,1)).text = '合并单元格'
        numbered(table.cell(1,0).paragraphs[0], 100, 0).add_run('编号')
        nested = table.cell(1,1).add_table(rows=1, cols=1)
        numbered(nested.cell(0,0).paragraphs[0], 100, 0).add_run('嵌套编号')
        document.save(self.source)
        original_file = self.source.read_bytes()
        before = read_parts(self.source)
        convert_docx_numbering(self.source, self.output)
        after = read_parts(self.output)
        self.assertEqual(self.source.read_bytes(), original_file)
        self.assertEqual(set(before),set(after))
        for name in before:
            if name != 'word/document.xml':
                self.assertEqual(before[name],after[name],name)
        originals, converted = paragraphs(self.source), paragraphs(self.output)
        for old, new in zip(originals, converted):
            old_content = [etree.tostring(e, method='c14n', exclusive=True) for e in old if e.tag != qn('w:pPr')]
            children = [e for e in new if e.tag != qn('w:pPr')]
            if old.find('w:pPr/w:numPr',NS) is not None:
                children = children[1:]
            self.assertEqual(old_content,[etree.tostring(e, method='c14n', exclusive=True) for e in children])
        old_root, new_root = etree.fromstring(before['word/document.xml']), etree.fromstring(after['word/document.xml'])
        for tag in ('tblPr','tblGrid','trPr','tcPr','drawing'):
            self.assertEqual([etree.tostring(e, method='c14n',exclusive=True) for e in old_root.iter(qn('w:'+tag))],
                             [etree.tostring(e, method='c14n',exclusive=True) for e in new_root.iter(qn('w:'+tag))])
        Document(self.output)  # package still readable by python-docx

    def test_headers_footers_notes_comments_textboxes_and_table_order(self):
        document = Document()
        add_list(document)
        numbered(document.add_paragraph('正文A'), 100, 0)
        host = document.add_paragraph('文本框宿主')
        textbox = element('txbxContent')
        inner = element('p')
        r = element('r')
        t = element('t'); t.text = '文本框'; r.append(t); inner.append(r)
        numbered(inner,100,0); textbox.append(inner)
        # Word's legacy text box container.
        v = 'urn:schemas-microsoft-com:vml'
        shape = etree.Element('{'+v+'}shape',nsmap={'v':v})
        vbox = etree.SubElement(shape,'{'+v+'}textbox'); vbox.append(textbox)
        pict = element('pict'); pict.append(shape)
        run = element('r'); run.append(pict); host._p.append(run)
        numbered(document.add_paragraph('正文B'), 100, 0)
        numbered(document.sections[0].header.paragraphs[0],100,0).add_run('页眉')
        numbered(document.sections[0].footer.paragraphs[0],100,0).add_run('页脚')
        document.save(self.source)
        parts = read_parts(self.source)
        types = etree.fromstring(parts['[Content_Types].xml'])
        rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
        updates = {}
        for name, singular in [('footnotes','footnote'),('endnotes','endnote'),('comments','comment')]:
            root = element(name)
            for identity in (1,2):
                note = element(singular,id=identity)
                for suffix in ('A','B'):
                    p = element('p'); numbered(p,100,0)
                    r = element('r'); t = element('t'); t.text = singular+str(identity)+suffix
                    r.append(t); p.append(r); note.append(p)
                root.append(note)
            updates['word/'+name+'.xml'] = etree.tostring(root)
            etree.SubElement(types,'{'+CT+'}Override',PartName='/word/'+name+'.xml',
                ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.'+name+'+xml')
            etree.SubElement(rels,'{'+REL+'}Relationship',Id='rId'+name,Type=OFFICE_REL+name,Target=name+'.xml')
        updates['[Content_Types].xml'] = etree.tostring(types)
        updates['word/_rels/document.xml.rels'] = etree.tostring(rels)
        rewrite_package(self.source,updates)
        report = convert_docx_numbering(self.source,self.output)
        self.assertEqual(report.converted,17)
        self.assertEqual(texts(self.output),['1.正文A','文本框宿主','1.文本框','2.正文B'])
        for name,singular in [('footnotes','footnote'),('endnotes','endnote'),('comments','comment')]:
            self.assertEqual(texts(self.output,'word/'+name+'.xml'),
                ['1.'+singular+'1A','2.'+singular+'1B','3.'+singular+'2A','4.'+singular+'2B'])
        self.assertEqual(texts(self.output,'word/header1.xml'),['1.页眉'])
        self.assertEqual(texts(self.output,'word/footer1.xml'),['1.页脚'])

    def test_jump_to_child_consumes_implicit_parent_numbers(self):
        document = Document()
        add_list(document, levels=[{'text':'%1.'},{'text':'%1.%2.'},{'text':'%1.%2.%3.'}])
        for level in (0,1,2,2,0,2,1,2):
            numbered(document.add_paragraph('X'),100,level)
        self.convert(document)
        self.assertEqual(texts(self.output),['1.X','1.1.X','1.1.1.X','1.1.2.X',
            '2.X','2.1.1.X','2.2.X','2.2.1.X'])

    def test_list_indentation_respects_inherited_and_direct_overrides(self):
        formatting_fixture(self.source)
        convert_docx_numbering(self.source,self.output)
        pp = paragraphs(self.output)
        expected = [('1800','600'),('1000','600'),('1800','100'),('1800','600')]
        for paragraph,(left,hanging) in zip(pp,expected):
            ind = paragraph.find('w:pPr/w:ind',NS)
            self.assertEqual(ind.get(qn('w:left')),left)
            self.assertEqual(ind.get(qn('w:hanging')),hanging)
        # WPS re-applies numbering layout after an ilvl-only paragraph override.
        ind = pp[5].find('w:pPr/w:ind',NS)
        self.assertEqual(ind.get(qn('w:left')),'720')
        self.assertEqual(ind.get(qn('w:hanging')),'360')

    def test_flush_left_stays_flush_left_including_table(self):
        document = Document()
        add_list(document)
        for p in (document.add_paragraph('正文顶格'),
                  document.add_table(rows=1,cols=1).cell(0,0).paragraphs[0]):
            numbered(p,100,0)
            p._p.get_or_add_pPr().append(element('ind',left=0,right=0,firstLine=0))
        self.convert(document)
        for p in paragraphs(self.output):
            ind = p.find('w:pPr/w:ind',NS)
            self.assertEqual(ind.get(qn('w:left')),'0')
            self.assertEqual(ind.get(qn('w:right')),'0')
            self.assertEqual(ind.get(qn('w:firstLine')),'0')
            self.assertIsNone(ind.get(qn('w:hanging')))

    def test_numbering_layout_reapplied_at_each_style_override(self):
        document = Document()
        add_list(document,levels=[{'text':'%1.'},{'text':'%1.%2.'}])
        base = document.styles.add_style('LayoutBase',WD_STYLE_TYPE.PARAGRAPH)
        base.element.append(element('pPr')); numbered(base.element,100,0)
        base.element.find(qn('w:pPr')).append(element('ind',left=1800,right=400,hanging=600))
        child = document.styles.add_style('LayoutChild',WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        child.element.append(element('pPr')); numbered(child.element,100,1)
        document.add_paragraph('继承基础样式',style=base)
        document.add_paragraph('子样式切换层级',style=child)
        self.convert(document)
        expected = [('1800','600'),('1440','360')]
        for p,(left,hanging) in zip(paragraphs(self.output),expected):
            ind = p.find('w:pPr/w:ind',NS)
            self.assertEqual(ind.get(qn('w:left')),left)
            self.assertEqual(ind.get(qn('w:hanging')),hanging)
            self.assertEqual(ind.get(qn('w:right')),'400')

    def test_character_unit_hanging_indent_survives_twip_first_line_override(self):
        document = Document()
        add_list(document)
        style = document.styles.add_style('CharacterIndented',WD_STYLE_TYPE.PARAGRAPH)
        style.element.append(element('pPr')); numbered(style.element,100,0)
        style.element.find(qn('w:pPr')).append(element('ind',leftChars=200,rightChars=100,hangingChars=100))
        p = document.add_paragraph('字符单位悬挂缩进',style=style)
        p._p.get_or_add_pPr().append(element('ind',firstLine=0))
        self.convert(document)
        ind = paragraphs(self.output)[0].find('w:pPr/w:ind',NS)
        # WPS keeps the original 2-char left / 1-char hanging indentation;
        # a later firstLine=0 in twips does not disable hangingChars.
        self.assertEqual(ind.get(qn('w:leftChars')),'200')
        self.assertEqual(ind.get(qn('w:rightChars')),'100')
        self.assertEqual(ind.get(qn('w:hangingChars')),'100')

    def test_zero_character_first_indent_keeps_list_twip_hanging(self):
        document = Document()
        add_list(document)
        p=numbered(document.add_paragraph('列表正文'),100,0)
        p._p.get_or_add_pPr().append(element('ind',firstLineChars=0))
        self.convert(document)
        ind=paragraphs(self.output)[0].find('w:pPr/w:ind',NS)
        self.assertEqual(ind.get(qn('w:hanging')),'360')
        self.assertEqual(ind.get(qn('w:firstLineChars')),'0')

    def test_japanese_unit_numbers_wps_boundaries(self):
        goldens={0:('〇','〇'),1:('一','壱'),2:('二','弐'),3:('三','参'),
                 10:('十','拾'),21:('二十一','弐拾壱'),100:('百','百'),
                 101:('百一','百壱'),110:('百十','百拾'),1001:('千一','阡壱'),
                 10010:('一万十','壱萬拾'),100000001:('一億一','壱億壱'),
                 2147483647:('二十一億四千七百四十八万三千六百四十七',
                             '弐拾壱億四阡七百四拾八萬参阡六百四拾七')}
        for value,(counting,legal) in goldens.items():
            self.assertEqual(format_number(value,'japaneseCounting'),counting)
            self.assertEqual(format_number(value,'japaneseLegal'),legal)
        for value in (-1,2147483648):
            with self.assertRaises(NumberingConversionError):
                format_number(value,'japaneseCounting')

    def test_japanese_mixed_outline_and_table(self):
        document=Document()
        add_list(document,100,[{'fmt':'japaneseCounting','text':'%1、'},
                               {'fmt':'decimal','text':'%1.%2.'}])
        numbered(document.add_paragraph('自动一级'),100,0)
        document.add_paragraph('二、手写不参与计数')
        numbered(document.add_paragraph('自动二级'),100,1)
        numbered(document.add_paragraph('下一一级'),100,0)
        table=document.add_table(rows=1,cols=1)
        numbered(table.cell(0,0).paragraphs[0],100,1).add_run('表格编号')
        add_list(document,101,[{'fmt':'japaneseLegal','start':1001,'text':'（%1）'}])
        numbered(document.add_paragraph('日式大写'),101,0)
        report=self.convert(document)
        self.assertEqual(report.converted,5)
        self.assertEqual(texts(self.output),['一、自动一级','二、手写不参与计数',
            '一.1.自动二级','二、下一一级','二.1.表格编号','（阡壱）日式大写'])

    def test_explicit_numbering_font_overrides_inherited_theme(self):
        document = Document()
        add_list(document)
        numbered(document.add_paragraph('编号字体必须保持一致'),100,0)
        self.convert(document)
        fonts = paragraphs(self.output)[0].find('w:r/w:rPr/w:rFonts',NS)
        self.assertEqual(fonts.get(qn('w:ascii')),'Arial')
        self.assertIsNone(fonts.get(qn('w:asciiTheme')))
        self.assertEqual(fonts.get(qn('w:eastAsia')),'宋体')
        self.assertIsNone(fonts.get(qn('w:eastAsiaTheme')))

    def test_tab_stops_and_first_line_indentation_preserved(self):
        document = Document()
        add_list(document,levels=[{'text':'%1.','suffix':'tab'}])
        style = document.styles.add_style('WithTabs',WD_STYLE_TYPE.PARAGRAPH)
        ppr = element('pPr')
        ppr.append(element('ind',left=800,hanging=200))
        tabs = element('tabs'); tabs.append(element('tab',val='right',pos=1800)); ppr.append(tabs)
        style.element.append(ppr)
        p = numbered(document.add_paragraph('A',style=style),100,0)
        direct = p._p.get_or_add_pPr()
        direct.append(element('ind',firstLine=100))
        tabs = element('tabs'); tabs.append(element('tab',val='clear',pos=1800)); direct.append(tabs)
        self.convert(document,preserve_suffix=True)
        converted = paragraphs(self.output)[0]
        ind = converted.find('w:pPr/w:ind',NS)
        self.assertEqual(ind.get(qn('w:firstLine')),'100')
        self.assertIsNone(ind.get(qn('w:hanging')))
        self.assertEqual([(e.get(qn('w:pos')),e.get(qn('w:val'))) for e in converted.findall('w:pPr/w:tabs/w:tab',NS)],
            [('720','left'),('1800','clear')])
        self.assertEqual(text(converted),'1.\tA')

    def test_section_restart_extension(self):
        from wfp_numbering import W15
        document = Document()
        abstract,_ = add_list(document)
        abstract.set('{'+W15+'}restartNumberingAfterBreak','1')
        numbered(document.add_paragraph('A'),100,0)
        numbered(document.add_paragraph('B'),100,0)
        document.add_section()
        numbered(document.add_paragraph('C'),100,0)
        self.convert(document)
        self.assertEqual(texts(self.output),['1.A','2.B','','1.C'])

    def test_known_office_format_disagreements_are_refused(self):
        for value,fmt in [(0,'upperRoman'),(1001,'upperRoman'),(0,'upperLetter'),
                          (1,'ideographLegalTraditional'),(1,'ideographZodiacTraditional')]:
            with self.subTest(fmt=fmt,value=value):
                with self.assertRaises(NumberingConversionError):
                    format_number(value,fmt)

    def test_interleaved_instances_sharing_abstract_are_refused(self):
        document = Document()
        add_list(document,100)
        _,instance = add_list(document,101,overrides={0:5})
        instance.find(qn('w:abstractNumId')).set(qn('w:val'),'100')
        for number in (100,101,100):
            numbered(document.add_paragraph('X'),number,0)
        document.save(self.source)
        with self.assertRaisesRegex(NumberingConversionError,'交错'):
            convert_docx_numbering(self.source,self.output)
        self.assertFalse(self.output.exists())
        report = convert_docx_numbering(self.source,self.output,strict=False)
        self.assertEqual(report.converted,0)
        self.assertEqual(texts(self.output),['X','X','X'])

    def test_expanded_samples_manual_content_and_all_numbers_removed(self):
        from wfp_numbering_samples import generate_scenarios
        for name,source in generate_scenarios(Path(self.directory.name)/'scenarios').items():
            if name == 'interleaved_lists_with_overrides':
                continue
            with self.subTest(scenario=name):
                output = source.with_name(source.stem+'_text.docx')
                original = paragraphs(source)
                active = [p.find('w:pPr/w:numPr/w:numId',NS) is not None for p in original]
                report = convert_docx_numbering(source,output)
                self.assertEqual(report.converted,sum(active))
                for was_numbered,before,after in zip(active,original,paragraphs(output)):
                    if was_numbered:
                        self.assertEqual(after.find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'0')
                    else:
                        self.assertEqual(text(before),text(after))

    def test_circular_definitions_are_refused(self):
        document = Document()
        style = document.styles.add_style('Cycle',WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = style
        document.add_paragraph('X',style=style)
        document.save(self.source)
        with self.assertRaisesRegex(NumberingConversionError,'循环'):
            convert_docx_numbering(self.source,self.output)
        self.assertFalse(self.output.exists())

    def test_idempotence_and_existing_output_protection(self):
        mixed_fixture(self.source)
        convert_docx_numbering(self.source,self.output)
        twice = self.output.with_name('twice.docx')
        report = convert_docx_numbering(self.output,twice)
        self.assertEqual(report.converted,0)
        self.assertEqual(read_parts(self.output),read_parts(twice))
        with self.assertRaises(FileExistsError):
            convert_docx_numbering(self.source,self.output)
        with self.assertRaises(ValueError):
            convert_docx_numbering(self.source,self.source,overwrite=True)

    def test_unsupported_formats_and_picture_bullets_fail_without_output(self):
        for fmt,picture in [('cardinalText',False),('bullet',True)]:
            with self.subTest(fmt=fmt):
                document = Document()
                abstract,_ = add_list(document,levels=[{'fmt':fmt,'text':'%1'}])
                if picture:
                    abstract.find(qn('w:lvl')).append(element('lvlPicBulletId',val=5))
                numbered(document.add_paragraph('X'),100,0)
                document.save(self.source)
                with self.assertRaises(NumberingConversionError):
                    convert_docx_numbering(self.source,self.output)
                self.assertFalse(self.output.exists())

    def test_numbering_fields_and_revisions_fail_without_output(self):
        document = Document()
        add_list(document)
        p = numbered(document.add_paragraph('X'),100,0)
        field = element('fldSimple',instr='LISTNUM LegalDefault \n')
        p._p.append(field)
        document.save(self.source)
        with self.assertRaisesRegex(NumberingConversionError,'LISTNUM'):
            convert_docx_numbering(self.source,self.output)
        p._p.remove(field)
        p._p.find('w:pPr/w:numPr',NS).append(element('numberingChange',id=1))
        document.save(self.source)
        with self.assertRaisesRegex(NumberingConversionError,'修订'):
            convert_docx_numbering(self.source,self.output)
        self.assertFalse(self.output.exists())

    def test_partial_mode_reports_and_preserves_unsupported_numbering(self):
        document = Document()
        add_list(document,100)
        add_list(document,101,[{'fmt':'cardinalText','text':'%1.'}])
        numbered(document.add_paragraph('A'),100,0)
        numbered(document.add_paragraph('B'),101,0)
        numbered(document.add_paragraph('C'),100,0)
        document.save(self.source)
        report = convert_docx_numbering(self.source,self.output,strict=False)
        self.assertEqual(report.converted,2)
        self.assertEqual(len(report.warnings),1)
        self.assertEqual(texts(self.output),['1.A','B','2.C'])
        self.assertEqual(paragraphs(self.output)[1].find('w:pPr/w:numPr/w:numId',NS).get(qn('w:val')),'101')

    def test_standard_character_bullets_and_none(self):
        document = Document()
        add_list(document,100,[{'fmt':'bullet','text':'●','suffix':'space'}])
        add_list(document,101,[{'fmt':'none','text':''}])
        numbered(document.add_paragraph('A'),100,0)
        numbered(document.add_paragraph('B'),101,0)
        self.convert(document)
        self.assertEqual(texts(self.output),['●A','B'])

    def test_chinese_numbers_verified_against_wps(self):
        cases = {0:'〇',1:'一',10:'十',11:'十一',20:'二十',101:'一百〇一',
                 1001:'一千〇一',10001:'一万〇一'}
        for fmt in ('chineseCounting','chineseCountingThousand'):
            for value,expected in cases.items():
                self.assertEqual(format_number(value,fmt),expected)
        self.assertEqual(format_number(10,'chineseLegalSimplified'),'壹拾')
        self.assertEqual(format_number(101,'chineseLegalSimplified'),'壹佰零壹')
        self.assertEqual(format_number(1,'ideographTraditional'),'甲')
        self.assertEqual(format_number(10,'ideographTraditional'),'癸')
        self.assertEqual(format_number(11,'ideographTraditional'),'11')


if __name__ == '__main__':
    unittest.main()
