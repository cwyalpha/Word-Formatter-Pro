"""Generate additional deterministic DOCX scenarios for WPS comparisons."""
from pathlib import Path
import random

from docx import Document
from docx.oxml.ns import qn

from wfp_numbering_tests import add_list, numbered, element


def generate_scenarios(directory):
    directory = Path(directory)
    directory.mkdir(parents=True,exist_ok=True)
    result = {}

    def save(name,document):
        path = directory / (name+'_input.docx')
        document.save(path)
        result[name] = path

    document = Document()
    add_list(document,levels=[{'text':'.'.join('%'+str(i+1) for i in range(depth+1))+'.'} for depth in range(4)])
    sequence = (0,1,2,3,3,2,3,1,2,3,0,1,2,3,1,2,3)
    for iteration in range(3):
        for index,level in enumerate(sequence):
            numbered(document.add_paragraph(f'全自动四级 {iteration}-{index}'),100,level)
    save('all_auto_decimal_outline',document)

    document = Document()
    add_list(document,levels=[{'fmt':'chineseCounting','text':'%1、'},
        {'fmt':'chineseCounting','text':'（%2）'},{'text':'%3.'},{'text':'（%4）'},
        {'fmt':'lowerLetter','text':'%5)'},{'fmt':'lowerRoman','text':'%6.'}])
    for iteration in range(4):
        for index,level in enumerate((0,1,2,3,4,5,5,4,5,3,4,5,2,3,4,5,1,2,3,4,5)):
            numbered(document.add_paragraph(f'全自动六种格式 {iteration}-{index}'),100,level)
    save('all_auto_chinese_outline',document)

    document = Document()
    formats = [('chineseCounting','%1、'),('chineseCounting','（%1）'),
               ('decimal','%1.'),('decimal','（%1）'),('lowerLetter','%1)'),('lowerRoman','%1.')]
    for number,(fmt,pattern) in enumerate(formats,100):
        add_list(document,number,[{'fmt':fmt,'text':pattern}])
    add_list(document,106,[{'text':'%1.'},{'text':'%1.%2.'},{'text':'%1.%2.%3.'}])
    manual = ['二、手写一级','（一）手写二级','2.手写三级','（2）手写四级',
              '1.2.1 手写复合编号','a) 手写字母','ii. 手写罗马数字','● 手写项目符号','无编号正文']
    generator = random.Random(20261004)
    for index in range(180):
        if index % 4 == 1 or index % 7 == 0:
            document.add_paragraph(generator.choice(manual)+f' #{index}')
        else:
            number = generator.randint(100,106)
            level = generator.randint(0,2) if number == 106 else 0
            numbered(document.add_paragraph(f'自动项目 #{index}'),number,level)
    save('mixed_numbering_forms',document)

    document = Document()
    add_list(document,levels=[{'text':'%1.'},{'text':'%1.%2.'},{'text':'%1.%2.%3.'}])
    numbered(document.add_paragraph('表格前一级'),100,0)
    table = document.add_table(rows=8,cols=3)
    for row in range(8):
        for column in range(3):
            cell = table.cell(row,column)
            p = cell.paragraphs[0]
            numbered(p,100,1 if column == 0 else 2)
            p.add_run(f'表格项目 {row}-{column}')
            cell.add_paragraph('9.9.9 表格手写编号')
            if (row,column) == (3,1):
                nested = cell.add_table(rows=1,cols=1)
                numbered(nested.cell(0,0).paragraphs[0],100,2).add_run('嵌套表格自动编号')
    numbered(document.add_paragraph('表格后三级续号'),100,2)
    numbered(document.add_paragraph('回到一级'),100,0)
    numbered(document.add_paragraph('第二级重启'),100,1)
    numbered(document.add_paragraph('第三级重启'),100,2)
    save('table_outline_continuation',document)

    document = Document()
    _,original = add_list(document,100,[{'text':'%1.'},{'text':'%1.%2.'},{'text':'%1.%2.%3.'}])
    for number,start in ((101,5),(102,20)):
        _,instance = add_list(document,number,overrides={0:start,1:3})
        instance.find(qn('w:abstractNumId')).set(qn('w:val'),'100')
    for iteration in range(4):
        for level in (0,1,2,2,1,2):
            for number in (100,101,102):
                numbered(document.add_paragraph(f'交错列表 {number}-{iteration}-{level}'),number,level)
                document.add_paragraph('8.8.8 手写，不能推进自动计数')
    save('interleaved_lists_with_overrides',document)

    document = Document()
    add_list(document,levels=[{'text':'.'.join('%'+str(i+1) for i in range(depth+1))+'）'} for depth in range(9)])
    for iteration in range(3):
        for index,level in enumerate(tuple(range(9))+(8,8,6,7,8,4,5,6,7,8,2,3,4,5,6,7,8)):
            numbered(document.add_paragraph(f'九级复合编号 {iteration}-{index}'),100,level)
    save('nine_level_outline',document)

    document = Document()
    number = 100
    for fmt in ('decimal','decimalZero','decimalFullWidth','decimalFullWidth2','ordinal','hex',
                'chineseCounting','chineseCountingThousand','chineseLegalSimplified'):
        for value in (1,2,9,19,21,99,109,110,1009,1010,1090,9999,10101,100000,100000001):
            add_list(document,number,[{'fmt':fmt,'text':'[%1]','start':value,'suffix':'space'}])
            numbered(document.add_paragraph(f'{fmt}:{value}'),number,0)
            number += 1
    save('number_formats_and_boundaries',document)

    document = Document()
    for number,bullet in enumerate(('●','•','■','◆','✓','►','—'),100):
        add_list(document,number,[{'fmt':'bullet','text':bullet,'suffix':'space'}])
        for index in range(3):
            numbered(document.add_paragraph(f'自动项目符号 {number}-{index}'),number,0)
            document.add_paragraph('● 手写符号保持原样')
    save('mixed_bullets',document)
    return result


if __name__ == '__main__':
    for name,path in generate_scenarios(Path(__file__).parent/'samples').items():
        print(name,path)
